from __future__ import annotations

import hashlib
import io
import os
import re
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from dotenv import load_dotenv
from pypdf import PdfReader
from pptx import Presentation


BACKEND_DIR = Path(__file__).resolve().parents[1]
load_dotenv(BACKEND_DIR / ".env")

SUPPORTED_DOCUMENT_SUFFIXES = {".pdf", ".pptx"}
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
CHUNK_SIZE = 900
CHUNK_OVERLAP = 100

STOPWORDS = {
    "一个", "一种", "以及", "通过", "进行", "可以", "这个", "这些", "我们", "你们", "他们",
    "主要", "相关", "内容", "问题", "方法", "使用", "实现", "需要", "基于", "目前", "当前",
    "the", "and", "for", "with", "from", "this", "that", "into", "are", "was", "were",
}


def normalize_text(value: str) -> str:
    value = value.replace("\x00", " ").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def _shape_texts(shape: Any) -> Iterable[str]:
    if getattr(shape, "has_text_frame", False):
        text = normalize_text(getattr(shape, "text", ""))
        if text:
            yield text
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                text = normalize_text(cell.text)
                if text:
                    yield text
    if hasattr(shape, "shapes"):
        for child in shape.shapes:
            yield from _shape_texts(child)


def extract_pptx(content: bytes) -> list[dict[str, Any]]:
    presentation = Presentation(io.BytesIO(content))
    pages: list[dict[str, Any]] = []
    for page_number, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            blocks.extend(_shape_texts(shape))
        blocks = list(dict.fromkeys(blocks))
        text = normalize_text("\n".join(blocks))
        title = ""
        if slide.shapes.title is not None:
            title = normalize_text(slide.shapes.title.text)
        if not title:
            title = next((line for line in text.splitlines() if line), f"第 {page_number} 页")
        pages.append({"pageNumber": page_number, "title": title[:180], "text": text})
    return pages


def extract_pdf(content: bytes) -> list[dict[str, Any]]:
    reader = PdfReader(io.BytesIO(content))
    pages: list[dict[str, Any]] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        title = next((line for line in text.splitlines() if line), f"第 {page_number} 页")
        pages.append({"pageNumber": page_number, "title": title[:180], "text": text})
    return pages


def extract_document(content: bytes, filename: str) -> tuple[str, list[dict[str, Any]]]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_SUFFIXES:
        raise ValueError("当前仅支持 .pptx 和 .pdf 文件")
    if not content:
        raise ValueError("上传文件为空")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("课件文件不能超过 25 MB")
    if suffix == ".pptx":
        return "pptx", extract_pptx(content)
    return "pdf", extract_pdf(content)


def split_chunks(text: str) -> list[str]:
    text = normalize_text(text)
    if not text:
        return []
    paragraphs = [item.strip() for item in re.split(r"\n+", text) if item.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > CHUNK_SIZE:
            if current:
                chunks.append(current)
                current = ""
            start = 0
            while start < len(paragraph):
                chunks.append(paragraph[start:start + CHUNK_SIZE])
                start += CHUNK_SIZE - CHUNK_OVERLAP
            continue
        candidate = f"{current}\n{paragraph}".strip()
        if current and len(candidate) > CHUNK_SIZE:
            chunks.append(current)
            current = f"{current[-CHUNK_OVERLAP:]}\n{paragraph}".strip()
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


def extract_keywords(pages: list[dict[str, Any]], limit: int = 60) -> list[dict[str, Any]]:
    counts: Counter[str] = Counter()
    display_names: dict[str, str] = {}
    for page in pages:
        candidates = [page["title"]]
        candidates.extend(line for line in page["text"].splitlines() if 2 <= len(line) <= 24)
        candidates.extend(re.findall(r"[A-Za-z][A-Za-z0-9+_.-]{2,30}", page["text"]))
        for candidate in candidates:
            cleaned = re.sub(r"^[\d一二三四五六七八九十、.．)）(（\-—•·\s]+", "", candidate).strip(" ：:，,。；;!?！？")
            normalized = cleaned.casefold()
            if not normalized or normalized in STOPWORDS or len(normalized) < 2 or len(normalized) > 40:
                continue
            if re.fullmatch(r"\d+(?:\.\d+)?", normalized):
                continue
            counts[normalized] += 2 if candidate == page["title"] else 1
            display_names.setdefault(normalized, cleaned)
    return [
        {"normalized": normalized, "name": display_names[normalized], "weight": weight}
        for normalized, weight in counts.most_common(limit)
    ]


def _term_id(normalized: str) -> str:
    digest = hashlib.sha1(normalized.encode("utf-8")).hexdigest()[:16]
    return f"cl_keyword_{digest}"


def build_document_graph(
    document_id: str,
    filename: str,
    document_type: str,
    pages: list[dict[str, Any]],
    session_id: str | None,
    created_at: float,
) -> dict[str, Any]:
    nodes: list[dict[str, Any]] = []
    edges: list[dict[str, Any]] = []
    document_node = {
        "id": document_id,
        "type": "Document",
        "label": filename,
        "properties": {
            "filename": filename,
            "documentType": document_type,
            "pageCount": len(pages),
            "namespace": "classloop",
            "createdAt": created_at,
        },
    }
    nodes.append(document_node)
    if session_id:
        session_node_id = f"cl_session_{session_id}"
        nodes.append({
            "id": session_node_id,
            "type": "Session",
            "label": "ClassLoop课堂",
            "properties": {"sessionId": session_id, "namespace": "classloop"},
        })
        edges.append({"id": f"edge_{session_node_id}_{document_id}", "source": session_node_id, "target": document_id, "type": "USES_DOCUMENT", "properties": {}})

    keyword_rows = extract_keywords(pages)
    for keyword in keyword_rows:
        nodes.append({
            "id": _term_id(keyword["normalized"]),
            "type": "Keyword",
            "label": keyword["name"],
            "properties": {**keyword, "namespace": "classloop", "source": "heuristic"},
        })

    previous_page_id: str | None = None
    for page in pages:
        page_number = page["pageNumber"]
        page_id = f"{document_id}_page_{page_number}"
        nodes.append({
            "id": page_id,
            "type": "Page",
            "label": page["title"],
            "properties": {
                "pageNumber": page_number,
                "title": page["title"],
                "text": page["text"],
                "charCount": len(page["text"]),
                "namespace": "classloop",
            },
        })
        edges.append({"id": f"edge_{document_id}_{page_id}", "source": document_id, "target": page_id, "type": "HAS_PAGE", "properties": {"order": page_number}})
        if previous_page_id:
            edges.append({"id": f"edge_{previous_page_id}_{page_id}", "source": previous_page_id, "target": page_id, "type": "NEXT_PAGE", "properties": {}})
        previous_page_id = page_id

        previous_chunk_id: str | None = None
        chunks = split_chunks(page["text"])
        page["chunks"] = chunks
        for chunk_index, chunk in enumerate(chunks, start=1):
            chunk_id = f"{page_id}_chunk_{chunk_index}"
            nodes.append({
                "id": chunk_id,
                "type": "Chunk",
                "label": f"第 {page_number} 页文本块 {chunk_index}",
                "properties": {"chunkIndex": chunk_index, "text": chunk, "charCount": len(chunk), "namespace": "classloop"},
            })
            edges.append({"id": f"edge_{page_id}_{chunk_id}", "source": page_id, "target": chunk_id, "type": "HAS_CHUNK", "properties": {"order": chunk_index}})
            if previous_chunk_id:
                edges.append({"id": f"edge_{previous_chunk_id}_{chunk_id}", "source": previous_chunk_id, "target": chunk_id, "type": "NEXT_CHUNK", "properties": {}})
            previous_chunk_id = chunk_id
            lowered = chunk.casefold()
            for keyword in keyword_rows:
                occurrences = lowered.count(keyword["normalized"])
                if occurrences:
                    keyword_id = _term_id(keyword["normalized"])
                    edges.append({
                        "id": f"edge_{chunk_id}_{keyword_id}",
                        "source": chunk_id,
                        "target": keyword_id,
                        "type": "MENTIONS",
                        "properties": {"count": occurrences},
                    })

    return {
        "schemaVersion": 1,
        "namespace": "classloop",
        "documentId": document_id,
        "nodes": nodes,
        "edges": edges,
        "stats": {
            "pages": len(pages),
            "chunks": sum(len(page.get("chunks", [])) for page in pages),
            "keywords": len(keyword_rows),
            "nodes": len(nodes),
            "edges": len(edges),
        },
    }


class ClassLoopNeo4jStore:
    NODE_LABELS = {
        "Document": "ClassLoopDocument",
        "Session": "ClassLoopSession",
        "Page": "ClassLoopPage",
        "Chunk": "ClassLoopChunk",
        "Keyword": "ClassLoopKeyword",
    }
    RELATION_TYPES = {"USES_DOCUMENT", "HAS_PAGE", "NEXT_PAGE", "HAS_CHUNK", "NEXT_CHUNK", "MENTIONS"}

    def __init__(self) -> None:
        self.uri = os.getenv("CLASSLOOP_NEO4J_URI", "").strip()
        self.user = os.getenv("CLASSLOOP_NEO4J_USER", "neo4j").strip()
        self.password = os.getenv("CLASSLOOP_NEO4J_PASSWORD", "")
        self.database = os.getenv("CLASSLOOP_NEO4J_DATABASE", "neo4j").strip() or "neo4j"

    @property
    def configured(self) -> bool:
        return bool(self.uri and self.password)

    def status(self) -> dict[str, Any]:
        if not self.configured:
            return {"configured": False, "connected": False, "database": self.database, "message": "尚未配置ClassLoop Neo4j连接"}
        try:
            from neo4j import GraphDatabase

            with GraphDatabase.driver(self.uri, auth=(self.user, self.password), connection_timeout=3) as driver:
                driver.verify_connectivity()
            return {"configured": True, "connected": True, "database": self.database, "message": "Neo4j连接正常"}
        except Exception as error:
            return {"configured": True, "connected": False, "database": self.database, "message": f"Neo4j连接失败：{error}"}

    def sync(self, graph: dict[str, Any]) -> dict[str, int]:
        if not self.configured:
            raise RuntimeError("尚未配置CLASSLOOP_NEO4J_URI和CLASSLOOP_NEO4J_PASSWORD")
        from neo4j import GraphDatabase

        def write_graph(tx: Any) -> None:
            for node in graph["nodes"]:
                label = self.NODE_LABELS[node["type"]]
                properties = {"label": node["label"], **node.get("properties", {})}
                tx.run(
                    f"MERGE (n:{label} {{id: $id, namespace: $namespace}}) SET n += $properties",
                    id=node["id"],
                    namespace="classloop",
                    properties=properties,
                )
            for edge in graph["edges"]:
                relation_type = edge["type"]
                if relation_type not in self.RELATION_TYPES:
                    continue
                tx.run(
                    f"MATCH (source {{id: $source, namespace: $namespace}}), "
                    f"(target {{id: $target, namespace: $namespace}}) "
                    f"MERGE (source)-[relation:{relation_type}]->(target) SET relation += $properties",
                    source=edge["source"],
                    target=edge["target"],
                    namespace="classloop",
                    properties=edge.get("properties", {}),
                )

        with GraphDatabase.driver(self.uri, auth=(self.user, self.password), connection_timeout=5) as driver:
            driver.verify_connectivity()
            with driver.session(database=self.database) as session:
                session.execute_write(write_graph)
        return {"nodes": len(graph["nodes"]), "edges": len(graph["edges"])}

    def delete_document(self, document_id: str) -> None:
        if not self.configured:
            return
        from neo4j import GraphDatabase

        query = """
        MATCH (document:ClassLoopDocument {id: $document_id, namespace: 'classloop'})
        OPTIONAL MATCH (document)-[:HAS_PAGE]->(page:ClassLoopPage)
        OPTIONAL MATCH (page)-[:HAS_CHUNK]->(chunk:ClassLoopChunk)
        DETACH DELETE chunk, page, document
        """
        cleanup_orphan_keywords = """
        MATCH (keyword:ClassLoopKeyword {namespace: 'classloop'})
        WHERE NOT ()-[:MENTIONS]->(keyword)
        DETACH DELETE keyword
        """
        with GraphDatabase.driver(self.uri, auth=(self.user, self.password), connection_timeout=5) as driver:
            with driver.session(database=self.database) as session:
                session.run(query, document_id=document_id).consume()
                session.run(cleanup_orphan_keywords).consume()

    def read_document(self, document_id: str) -> dict[str, Any]:
        """Read one ClassLoop document subgraph directly from Neo4j."""
        if not self.configured:
            raise RuntimeError("尚未配置CLASSLOOP_NEO4J_URI和CLASSLOOP_NEO4J_PASSWORD")
        from neo4j import GraphDatabase

        node_query = """
        MATCH (document:ClassLoopDocument {id: $document_id, namespace: 'classloop'})
        OPTIONAL MATCH (class_session:ClassLoopSession {namespace: 'classloop'})-[:USES_DOCUMENT]->(document)
        OPTIONAL MATCH (document)-[:HAS_PAGE]->(page:ClassLoopPage {namespace: 'classloop'})
        OPTIONAL MATCH (page)-[:HAS_CHUNK]->(chunk:ClassLoopChunk {namespace: 'classloop'})
        OPTIONAL MATCH (chunk)-[:MENTIONS]->(keyword:ClassLoopKeyword {namespace: 'classloop'})
        WITH document, collect(DISTINCT class_session) AS class_sessions,
             collect(DISTINCT page) AS pages, collect(DISTINCT chunk) AS chunks,
             collect(DISTINCT keyword) AS keywords
        WITH [document] + class_sessions + pages + chunks + keywords AS candidates
        UNWIND candidates AS node
        WITH DISTINCT node
        WHERE node IS NOT NULL
        RETURN node.id AS id, labels(node)[0] AS node_type, properties(node) AS properties
        """
        relationship_query = """
        MATCH (document:ClassLoopDocument {id: $document_id, namespace: 'classloop'})
        OPTIONAL MATCH (class_session:ClassLoopSession {namespace: 'classloop'})-[:USES_DOCUMENT]->(document)
        OPTIONAL MATCH (document)-[:HAS_PAGE]->(page:ClassLoopPage {namespace: 'classloop'})
        OPTIONAL MATCH (page)-[:HAS_CHUNK]->(chunk:ClassLoopChunk {namespace: 'classloop'})
        OPTIONAL MATCH (chunk)-[:MENTIONS]->(keyword:ClassLoopKeyword {namespace: 'classloop'})
        WITH document, collect(DISTINCT class_session) AS class_sessions,
             collect(DISTINCT page) AS pages, collect(DISTINCT chunk) AS chunks,
             collect(DISTINCT keyword) AS keywords
        WITH [document] + class_sessions + pages + chunks + keywords AS candidates
        UNWIND candidates AS node
        WITH collect(DISTINCT node.id) AS ids
        MATCH (source {namespace: 'classloop'})-[relationship]->(target {namespace: 'classloop'})
        WHERE source.id IN ids AND target.id IN ids
              AND type(relationship) IN $relationship_types
        RETURN source.id AS source, target.id AS target, type(relationship) AS relationship_type,
               properties(relationship) AS properties
        """
        reverse_labels = {value: key for key, value in self.NODE_LABELS.items()}
        with GraphDatabase.driver(self.uri, auth=(self.user, self.password), connection_timeout=5) as driver:
            driver.verify_connectivity()
            with driver.session(database=self.database) as session:
                node_records = list(session.run(node_query, document_id=document_id))
                if not node_records:
                    raise LookupError("Neo4j中未找到该课件图谱")
                relationship_records = list(session.run(
                    relationship_query,
                    document_id=document_id,
                    relationship_types=sorted(self.RELATION_TYPES),
                ))

        nodes: list[dict[str, Any]] = []
        for record in node_records:
            properties = dict(record["properties"] or {})
            properties.pop("id", None)
            label = properties.pop("label", record["id"])
            nodes.append({
                "id": record["id"],
                "type": reverse_labels.get(record["node_type"], record["node_type"]),
                "label": label,
                "properties": properties,
            })
        edges = [{
            "id": f"neo4j_{record['source']}_{record['relationship_type']}_{record['target']}",
            "source": record["source"],
            "target": record["target"],
            "type": record["relationship_type"],
            "properties": dict(record["properties"] or {}),
        } for record in relationship_records]
        return {"namespace": "classloop", "documentId": document_id, "nodes": nodes, "edges": edges}


neo4j_store = ClassLoopNeo4jStore()
